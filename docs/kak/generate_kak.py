#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generator KAK (Kerangka Acuan Kerja)
Pengembangan Aplikasi Dek-Oke (Deklarasi Konflik Kepentingan)
Dinas Komunikasi dan Informatika Kabupaten Aceh Tengah - TA 2026

Menghasilkan: KAK-Dek-Oke-TA2026.docx
Cara pakai:   python generate_kak.py
Butuh:        pip install python-docx
"""

import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# =====================================================================
# PARAMETER DOKUMEN — ubah di sini bila ada perubahan data
# =====================================================================
PARAM = {
    "judul": "PENGEMBANGAN APLIKASI DEKLARASI KONFLIK KEPENTINGAN (DEK-OKE) BERBASIS WEB",
    "instansi": "PEMERINTAH KABUPATEN ACEH TENGAH",
    "opd_pelaksana": "Dinas Komunikasi dan Informatika Kabupaten Aceh Tengah",
    "tahun": "2026",
    "pagu": "Rp 10.000.000,-",
    "pagu_kata": "Sepuluh Juta Rupiah",
    "kegiatan": "Layanan e-Government",
    "sub_kegiatan": "Koordinasi dan Fasilitasi Penyelenggaraan Kabupaten atau Kota Cerdas",
    "sumber_dana": "APBK Kabupaten Aceh Tengah Tahun Anggaran 2026",
}

FONT = "Times New Roman"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "KAK-Dek-Oke-TA2026.docx")

doc = Document()

# ---------------------------------------------------------------------
# Gaya & ukuran halaman
# ---------------------------------------------------------------------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)          # A4
sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.2)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(12)
rpr = normal.element.get_or_add_rPr()
rfonts = rpr.get_or_add_rFonts()
for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
    rfonts.set(qn(attr), FONT)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for sname in ("List Bullet", "List Number", "Table Grid"):
    try:
        st = doc.styles[sname]
        st.font.name = FONT
        st.font.size = Pt(11)
    except KeyError:
        pass


# ---------------------------------------------------------------------
# Helper
# ---------------------------------------------------------------------
def _fix_font(paragraph, size=None, bold=None):
    for run in paragraph.runs:
        run.font.name = FONT
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold
        r = run._element.get_or_add_rPr()
        rf = r.get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(attr), FONT)
    return paragraph


def p(text="", size=12, bold=False, italic=False, align="just", space_after=6,
      space_before=0, indent_left=None, first_line_indent=None, line=1.15):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold, run.italic = bold, italic
    par.alignment = {
        "just": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = par.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(space_after), Pt(space_before), line
    if indent_left:
        pf.left_indent = Cm(indent_left)
    if first_line_indent:
        pf.first_line_indent = Cm(first_line_indent)
    _fix_font(par, size=size, bold=bold)
    if italic:
        for r in par.runs:
            r.italic = True
    return par


def p_mixed(parts, size=12, align="just", space_after=6, indent_left=None, line=1.15):
    """parts = [(teks, bold), ...]"""
    par = doc.add_paragraph()
    for text, bold in parts:
        run = par.add_run(text)
        run.bold = bold
    par.alignment = {"just": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                     "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf = par.paragraph_format
    pf.space_after, pf.line_spacing = Pt(space_after), line
    if indent_left:
        pf.left_indent = Cm(indent_left)
    _fix_font(par, size=size)
    return par


def bab(roman, judul):
    """Judul bab: I. PENDAHULUAN"""
    par = doc.add_paragraph()
    par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(14), Pt(6)
    par.paragraph_format.keep_with_next = True
    par.add_run(f"{roman}. {judul.upper()}").bold = True
    _fix_font(par, size=12, bold=True)
    return par


def sub(huruf_atau_angka, judul, space_before=8):
    """Sub-bab: A. Latar Belakang / 1. Tahapan"""
    par = doc.add_paragraph()
    par.paragraph_format.space_before, par.paragraph_format.space_after = Pt(space_before), Pt(4)
    par.paragraph_format.keep_with_next = True
    par.add_run(f"{huruf_atau_angka}. {judul}").bold = True
    _fix_font(par, size=12, bold=True)
    return par


def bullets(items, num=False, size=11):
    for it in items:
        par = doc.add_paragraph(style="List Number" if num else "List Bullet")
        if isinstance(it, tuple):
            par.add_run(it[0]).bold = True
            par.add_run(it[1])
        else:
            par.add_run(it)
        par.paragraph_format.space_after = Pt(2)
        par.paragraph_format.line_spacing = 1.15
        _fix_font(par, size=size)


def shade(cell, fill="D9D9D9"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def table(headers, rows, widths, size=10.5, header_fill="D9D9D9", align_center_cols=(),
          align_right_cols=(), repeat_header=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        par = hdr[i].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_after = Pt(2)
        par.add_run(h).bold = True
        _fix_font(par, size=size, bold=True)
        shade(hdr[i], header_fill)
    if repeat_header:  # ulangi baris kepala bila tabel berlanjut ke halaman berikutnya
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing = 1.0
            if i in align_center_cols:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i in align_right_cols:
                par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            par.add_run(str(val))
            _fix_font(par, size=size)
    # lebar kolom
    for r in t.rows:
        for i, w in enumerate(widths):
            r.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def tabel_kv(pairs, width_kiri=4.5, width_kanan=10.5):
    t = doc.add_table(rows=0, cols=3)
    t.style = "Table Grid"
    t.autofit = False
    for k, v in pairs:
        cells = t.add_row().cells
        for i, txt in enumerate((k, ":", v)):
            cells[i].text = ""
            par = cells[i].paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing = 1.0
            par.add_run(txt)
            if i < 2:
                par.runs[0].bold = True
            _fix_font(par, size=11)
        cells[0].width, cells[1].width, cells[2].width = Cm(width_kiri), Cm(0.5), Cm(width_kanan)
    return t


def ttd_kolom(kiri, kanan, tinggi_spasi=4):
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    for idx, blok in enumerate((kiri, kanan)):
        cell = t.rows[0].cells[idx]
        cell.width = Cm(7.75)
        cell.text = ""
        first = True
        for line in blok:
            par = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.line_spacing = 1.1
            par.add_run(line[0]).bold = line[1]
            _fix_font(par, size=12)
        for _ in range(tinggi_spasi):
            par = cell.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_after = Pt(2)
    return t


def nomor_halaman(paragraph):
    run = paragraph.add_run()
    for el, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "end")):
        node = OxmlElement(el)
        if el == "w:fldChar":
            node.set(qn("w:fldCharType"), attr)
        else:
            node.set(qn("xml:space"), "preserve")
            node.text = "PAGE"
        run._r.append(node)
    _fix_font(paragraph, size=9)


def catatan(text):
    """Paragraf catatan miring berukuran kecil."""
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.5)
    par.paragraph_format.space_after = Pt(8)
    run = par.add_run(text)
    run.italic = True
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _fix_font(par, size=10.5)
    for r in par.runs:
        r.italic = True
    return par


# ---------------------------------------------------------------------
# Kepala & kaki halaman
# ---------------------------------------------------------------------
hdr_par = sec.header.paragraphs[0]
hdr_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
hr = hdr_par.add_run("KAK Pengembangan Aplikasi Dek-Oke — Diskominfo Kabupaten Aceh Tengah TA 2026")
hr.italic = True
_fix_font(hdr_par, size=9)
for r in hdr_par.runs:
    r.italic = True

ftr_par = sec.footer.paragraphs[0]
ftr_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
_fix_font(ftr_par, size=9)
ftr_par.add_run("Halaman ")
nomor_halaman(ftr_par)

# =====================================================================
# HALAMAN JUDUL
# =====================================================================
_logo = os.path.join(os.path.dirname(OUT), "..", "..", "assets", "logo-aceh-tengah.png")
if os.path.exists(_logo):
    _par = doc.add_paragraph()
    _par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _par.paragraph_format.space_before = Pt(0)
    _par.paragraph_format.space_after = Pt(2)
    _par.add_run().add_picture(_logo, width=Cm(2.5))

p("PEMERINTAH KABUPATEN ACEH TENGAH", size=12, bold=True, align="center", space_after=1)
p(PARAM["opd_pelaksana"].upper(), size=12, bold=True, align="center", space_after=16)

p("KERANGKA ACUAN KERJA (KAK)", size=16, bold=True, align="center", space_before=6, space_after=2)
p("(TERMS OF REFERENCE)", size=11, italic=True, align="center", space_after=20)
p(PARAM["judul"], size=14, bold=True, align="center", space_after=6)
p("DI LINGKUNGAN " + PARAM["instansi"], size=13, bold=True, align="center", space_after=6)
p(f"TAHUN ANGGARAN {PARAM['tahun']}", size=13, bold=True, align="center", space_after=26)

tabel_kv([
    ("Program", "………………………… (sesuai DPA-SKPD)"),
    ("Kegiatan", PARAM["kegiatan"]),
    ("Sub Kegiatan", PARAM["sub_kegiatan"]),
    ("Kode Rekening", "…………………………"),
    ("Pagu Anggaran", f"{PARAM['pagu']} ({PARAM['pagu_kata']})"),
    ("Sumber Dana", PARAM["sumber_dana"]),
    ("Pelaksana Anggaran", PARAM["opd_pelaksana"]),
    ("Penanggung Jawab", "Pejabat Pelaksana Teknis Kegiatan (PPTK)"),
    ("Pengguna Layanan", "Inspektorat dan seluruh Perangkat Daerah Kabupaten Aceh Tengah"),
    ("Jangka Waktu", "2 (dua) bulan sejak diterbitkannya Surat Pesanan/Kontrak"),
])

p()
p(PARAM["instansi"], size=12, bold=True, align="center", space_before=30, space_after=2)
p("TAHUN " + PARAM["tahun"], size=12, bold=True, align="center")

doc.add_page_break()

# =====================================================================
# LEMBAR PENGESAHAN
# =====================================================================
p("LEMBAR PENGESAHAN", size=13, bold=True, align="center", space_after=14)
p("Kerangka Acuan Kerja (KAK) dengan judul:", align="center", space_after=6)
p(PARAM["judul"], bold=True, align="center", space_after=6)
p("yang disusun sebagai dasar pelaksanaan Sub Kegiatan "
  f"{PARAM['sub_kegiatan']} pada {PARAM['opd_pelaksana']} "
  f"Tahun Anggaran {PARAM['tahun']} dengan nilai pagu {PARAM['pagu']} ({PARAM['pagu_kata']}), "
  "disahkan untuk dipergunakan sebagaimana mestinya.", align="center", space_after=24)

p("Takengon, ………………………… " + PARAM["tahun"], align="right", space_after=18)

ttd_kolom(
    [
        ("Dibuat oleh,", False),
        ("Pejabat Pelaksana Teknis Kegiatan (PPTK)", False),
        ("Sub Kegiatan Koordinasi dan Fasilitasi", False),
        ("Penyelenggaraan Kabupaten/Kota Cerdas", False),
        ("", False),
        ("", False),
        ("……………………………………", True),
        ("NIP. ……………………………………", False),
    ],
    [
        ("Disetujui oleh,", False),
        ("Inspektur Kabupaten Aceh Tengah", False),
        ("", False),
        ("", False),
        ("", False),
        ("", False),
        ("……………………………………", True),
        ("NIP. ……………………………………", False),
    ],
)

catatan("Catatan: nama, jabatan, NIP, dan tanggal pengesahan agar diisi/disesuaikan dengan "
        "struktur organisasi serta ketentuan yang berlaku pada Pemerintah Kabupaten Aceh Tengah "
        "(misalnya penandatangan dapat disesuaikan menjadi Kepala Dinas Komunikasi dan Informatika "
        "selaku Pengguna Anggaran/Kuasa Pengguna Anggaran).")

doc.add_page_break()

# =====================================================================
# I. PENDAHULUAN
# =====================================================================
bab("I", "Pendahuluan")

sub("A", "Latar Belakang")
p("Konflik kepentingan adalah kondisi Pejabat Pemerintahan yang memiliki kepentingan pribadi "
  "untuk menguntungkan diri sendiri dan/atau orang lain dalam penggunaan wewenang, sehingga dapat "
  "memengaruhi netralitas dan kualitas keputusan dan/atau tindakan yang dibuat dan/atau dilakukannya. "
  "Konflik kepentingan yang tidak dikelola dengan baik merupakan salah satu pemicu utama praktik "
  "korupsi, kolusi, dan nepotisme (KKN). Oleh karena itu, pengelolaan konflik kepentingan merupakan "
  "bagian penting dari upaya pencegahan korupsi dan pembangunan birokrasi yang bersih dan berintegritas.")
p("Peraturan Menteri Pendayagunaan Aparatur Negara dan Reformasi Birokrasi Nomor 17 Tahun 2024 "
  "tentang Pengelolaan Konflik Kepentingan (yang menggantikan Permen PANRB Nomor 37 Tahun 2012) "
  "mewajibkan setiap Instansi Pemerintah membangun dan melaksanakan Sistem Pengelolaan Konflik "
  "Kepentingan yang meliputi: (a) pencatatan Daftar Kepentingan Pribadi; (b) Deklarasi Konflik "
  "Kepentingan; (c) pengendalian/penanganan atas konflik kepentingan; (d) pengendalian melalui masa "
  "tunggu (cooling period); serta (e) pelatihan dan konsultasi. Pencatatan Daftar Kepentingan Pribadi "
  "dilaksanakan secara berkala, sekurang-kurangnya 1 (satu) kali dalam 1 (satu) tahun, dan dimuat "
  "dalam sistem informasi instansi.")
p("Praktik pengelolaan konflik kepentingan di banyak perangkat daerah saat ini masih dilakukan secara "
  "manual (formulir kertas atau lembar kerja elektronik yang tidak terintegrasi). Kondisi tersebut "
  "menimbulkan berbagai kendala, antara lain: risiko dokumen hilang atau rusak; proses rekapitulasi "
  "yang lambat dan rawan salah hitung; tidak adanya pemantauan status penanganan secara terpusat; "
  "tanda tangan basah yang mengharuskan kehadiran fisik; kesulitan penelusuran (audit trail); serta "
  "pemborosan kertas dan ruang penyimpanan. Kondisi ini juga menyulitkan Inspektorat selaku pengelola "
  "dan pengawas dalam memetakan risiko konflik kepentingan secara cepat dan akurat.")
p("Sejalan dengan kebijakan Sistem Pemerintahan Berbasis Elektronik (SPBE) dan program Kabupaten/Kota "
  "Cerdas (Smart City/Smart Regency), serta dalam rangka percepatan digitalisasi layanan administrasi "
  "pemerintahan, Pemerintah Kabupaten Aceh Tengah telah mengembangkan aplikasi berbasis web bernama "
  "Dek-Oke (Deklarasi Konflik Kepentingan). Aplikasi ini menyediakan pengisian Formulir Daftar "
  "Kepentingan Pribadi dan Formulir Deklarasi Konflik Kepentingan secara daring, dilengkapi tanda "
  "tangan digital, penyimpanan basis data, serta kanal pengelolaan oleh administrator.")
p("Agar aplikasi tersebut dapat digunakan secara luas, aman, andal, dan berkelanjutan, diperlukan "
  "kegiatan pengembangan lanjutan yang mencakup pemantapan dan penyesuaian fitur, penguatan keamanan "
  "dan pelindungan data pribadi, penyediaan infrastruktur produksi, pelatihan dan sosialisasi, serta "
  "pemeliharaan teknis. Sehubungan dengan hal tersebut, disusun Kerangka Acuan Kerja (KAK) ini sebagai "
  "acuan pelaksanaan, pengendalian, dan pertanggungjawaban kegiatan pada Tahun Anggaran "
  f"{PARAM['tahun']}.")

sub("B", "Gambaran Umum")
p("Aplikasi Dek-Oke dirancang sebagai aplikasi berbasis web yang dapat diakses melalui peramban "
  "(browser) tanpa perlu pemasangan perangkat lunak pada komputer pengguna, dengan gambaran umum "
  "sebagai berikut:")
bullets([
    ("Cakupan formulir: ", "Formulir Daftar Kepentingan Pribadi (identitas pegawai dan Bagian A sampai "
     "dengan F) serta Formulir Deklarasi Konflik Kepentingan (identitas, atasan pejabat, jenis dan "
     "sumber konflik, uraian, serta pengendalian yang disarankan)."),
    ("Tanda tangan digital: ", "kanvas tanda tangan yang dapat ditulis menggunakan mouse atau layar "
     "sentuh, dengan fungsi undo dan hapus."),
    ("Penyimpanan data: ", "basis data relasional (PostgreSQL/MySQL) beserta penyimpanan sementara di "
     "peramban agar isian tidak hilang sebelum dikirim."),
    ("Halaman administrator: ", "rekapitulasi data, lembar isian read-only yang tampil persis seperti "
     "formulir asli (lengkap dengan tanda tangan) dan siap dicetak, pengubahan status dan catatan, "
     "serta ekspor data (CSV)."),
    ("Keluaran dokumen: ", "cetak/ekspor PDF ukuran A4 dengan tata letak menyerupai formulir asli."),
    ("Kesiapan infrastruktur: ", "dapat ditempatkan pada layanan awan maupun server milik Pemerintah "
     "Daerah."),
])
p("Meskipun versi awal telah tersedia, aplikasi belum sepenuhnya siap digunakan sebagai layanan "
  "produksi berskala instansi. Beberapa hal yang masih diperlukan antara lain: penyempurnaan alur "
  "pengisian dan validasi data, penyesuaian format formulir dengan ketentuan terbaru, penguatan "
  "keamanan aplikasi dan pelindungan data pribadi, penyediaan lingkungan produksi (hosting, basis "
  "data, domain, dan sertifikat keamanan), penyusunan dokumentasi, serta pelatihan bagi pengguna dan "
  "administrator. Pekerjaan dimaksud akan dilaksanakan oleh penyedia jasa pengembangan perangkat "
  "lunak melalui mekanisme pengadaan barang/jasa sesuai ketentuan yang berlaku.")

sub("C", "Maksud dan Tujuan")
p_mixed([("Maksud. ", True),
         ("KAK ini disusun sebagai acuan teknis, administratif, dan anggaran bagi pelaksanaan "
          "pengembangan aplikasi Dek-Oke pada Tahun Anggaran " + PARAM["tahun"] +
          ", sehingga kegiatan dapat dilaksanakan secara terarah, terukur, dan dapat "
          "dipertanggungjawabkan.", False)])
p_mixed([("Tujuan. ", True), ("Kegiatan ini bertujuan untuk:", False)])
bullets([
    "mendigitalisasi pencatatan Daftar Kepentingan Pribadi dan Deklarasi Konflik Kepentingan di "
    "lingkungan Pemerintah Kabupaten Aceh Tengah;",
    "menyediakan sarana pelaporan konflik kepentingan yang cepat, mudah, aman, dan tertelusur bagi "
    "seluruh ASN dan Pejabat Pemerintahan;",
    "mempercepat dan mempermudah Inspektorat dalam melakukan rekapitulasi, analisis risiko, "
    "pemantauan status, dan pengendalian konflik kepentingan;",
    "menghasilkan dokumen deklarasi elektronik yang sah secara administratif dan siap cetak sebagai "
    "arsip;",
    "mendukung pencapaian nilai Sistem Pemerintahan Berbasis Elektronik (SPBE), Zona Integritas "
    "menuju Wilayah Bebas dari Korupsi (ZI-WBK), serta program Kabupaten/Kota Cerdas;",
    "menyediakan data dan informasi yang akurat sebagai dasar pengambilan kebijakan pengelolaan "
    "konflik kepentingan oleh pimpinan.",
], num=True)

sub("D", "Sasaran")
bullets([
    "Tersedianya aplikasi Dek-Oke versi produksi yang berjalan pada infrastruktur yang memadai dan "
    "dapat diakses oleh seluruh perangkat daerah.",
    "Terlaksananya pengisian Daftar Kepentingan Pribadi secara berkala dan Deklarasi Konflik "
    "Kepentingan secara daring oleh ASN dan Pejabat Pemerintahan di lingkungan Pemerintah Kabupaten "
    "Aceh Tengah (jumlah pengguna sasaran: …………… orang, sesuai data Badan Kepegawaian dan "
    "Pengembangan Sumber Daya Manusia Kabupaten Aceh Tengah).",
    "Tersedianya kanal pengelolaan dan pemantauan (dashboard) bagi administrator pada Inspektorat "
    "Kabupaten Aceh Tengah.",
    "Terpenuhinya aspek keamanan aplikasi dan pelindungan data pribadi sesuai ketentuan peraturan "
    "perundang-undangan.",
    "Tersedianya dokumentasi, manual penggunaan, dan pelatihan bagi pengguna serta administrator.",
], num=True)

sub("E", "Dasar Hukum")
bullets([
    "Undang-Undang Nomor 28 Tahun 1999 tentang Penyelenggaraan Negara yang Bersih dan Bebas dari "
    "Korupsi, Kolusi, dan Nepotisme;",
    "Undang-Undang Nomor 31 Tahun 1999 tentang Pemberantasan Tindak Pidana Korupsi sebagaimana telah "
    "diubah dengan Undang-Undang Nomor 20 Tahun 2001;",
    "Undang-Undang Nomor 30 Tahun 2014 tentang Administrasi Pemerintahan (khususnya Pasal 42 sampai "
    "dengan Pasal 45 mengenai larangan dan pelaporan konflik kepentingan);",
    "Undang-Undang Nomor 5 Tahun 2014 tentang Aparatur Sipil Negara sebagaimana telah diubah dengan "
    "Undang-Undang Nomor 20 Tahun 2023;",
    "Undang-Undang Nomor 27 Tahun 2022 tentang Pelindungan Data Pribadi;",
    "Peraturan Pemerintah Nomor 94 Tahun 2021 tentang Disiplin Pegawai Negeri Sipil;",
    "Peraturan Pemerintah Nomor 12 Tahun 2019 tentang Pengelolaan Keuangan Daerah;",
    "Peraturan Presiden Nomor 95 Tahun 2018 tentang Sistem Pemerintahan Berbasis Elektronik "
    "sebagaimana dilengkapi dengan Peraturan Presiden Nomor 132 Tahun 2022 tentang Arsitektur Sistem "
    "Pemerintahan Berbasis Elektronik Nasional;",
    "Peraturan Menteri Pendayagunaan Aparatur Negara dan Reformasi Birokrasi Nomor 17 Tahun 2024 "
    "tentang Pengelolaan Konflik Kepentingan;",
    "Peraturan Menteri Pendayagunaan Aparatur Negara dan Reformasi Birokrasi Nomor 59 Tahun 2020 "
    "tentang Pemantauan dan Evaluasi Sistem Pemerintahan Berbasis Elektronik;",
    "Peraturan Menteri Dalam Negeri Nomor 77 Tahun 2020 tentang Pedoman Teknis Pengelolaan Keuangan "
    "Daerah;",
    "Peraturan Bupati Aceh Tengah Nomor ……… Tahun " + PARAM["tahun"] + " tentang Penjabaran Anggaran "
    "Pendapatan dan Belanja Kabupaten (APBK) Tahun Anggaran " + PARAM["tahun"] + ";",
    "Dokumen Pelaksanaan Anggaran Satuan Kerja Perangkat Daerah (DPA-SKPD) Dinas Komunikasi dan "
    "Informatika Kabupaten Aceh Tengah Tahun Anggaran " + PARAM["tahun"] + ".",
], num=True)

# =====================================================================
# II. PENERIMA MANFAAT
# =====================================================================
bab("II", "Penerima Manfaat")
p("Penerima manfaat kegiatan pengembangan aplikasi Dek-Oke adalah sebagai berikut:")
table(
    ["No.", "Penerima Manfaat", "Manfaat yang Diterima"],
    [
        ["1", "Inspektorat Kabupaten Aceh Tengah",
         "Memperoleh sarana pengelolaan, pemantauan, dan analisis risiko konflik kepentingan secara "
         "terpusat, cepat, dan tertelusur, serta kemudahan rekapitulasi dan pelaporan."],
        ["2", "Seluruh ASN dan Pejabat Pemerintahan Kabupaten Aceh Tengah",
         "Memperoleh kemudahan mengisi Daftar Kepentingan Pribadi dan mendeklarasikan konflik "
         "kepentingan secara daring, kapan saja, tanpa kehadiran fisik, serta memperoleh bukti "
         "pengisian dan dokumen siap cetak."],
        ["3", "Perangkat Daerah",
         "Memperoleh dokumen deklarasi yang tertata dan terarsip secara elektronik serta mendukung "
         "pemenuhan indikator pengelolaan konflik kepentingan."],
        ["4", "Dinas Komunikasi dan Informatika Kabupaten Aceh Tengah",
         "Bertambahnya aplikasi layanan e-Government yang terkelola, terdokumentasi, dan mendukung "
         "capaian indikator SPBE serta Kabupaten/Kota Cerdas."],
        ["5", "Pimpinan Daerah dan Masyarakat",
         "Meningkatnya transparansi, akuntabilitas, dan kepercayaan publik terhadap penyelenggaraan "
         "pemerintahan yang bersih dan bebas dari konflik kepentingan."],
    ],
    widths=[1.1, 4.6, 9.3],
    align_center_cols=(0,),
)

# =====================================================================
# III. HASIL YANG DIHARAPKAN
# =====================================================================
bab("III", "Hasil yang Diharapkan")
p("Kegiatan ini diharapkan menghasilkan keluaran (output) sebagai berikut:")
table(
    ["No.", "Keluaran", "Volume", "Keterangan"],
    [
        ["1", "Aplikasi Dek-Oke versi produksi yang aktif dan dapat diakses oleh pengguna internal "
              "melalui alamat/domain yang ditetapkan", "1 aplikasi",
         "Berjalan pada lingkungan produksi dengan domain dan sertifikat keamanan (HTTPS)."],
        ["2", "Modul Formulir Daftar Kepentingan Pribadi (Bagian A–F) dan Formulir Deklarasi Konflik "
              "Kepentingan yang disempurnakan", "2 modul",
         "Termasuk validasi isian, penyimpanan sementara, tanda tangan digital, dan ekspor PDF A4."],
        ["3", "Modul administrator (rekapitulasi, detail lembar isian, status, catatan, ekspor data)",
         "1 modul", "Dilengkapi autentikasi dan pembatasan hak akses."],
        ["4", "Basis data pengisian dan deklarasi konflik kepentingan beserta prosedur pencadangan",
         "1 basis data", "Skema tabel terpasang dan berfungsi pada server produksi."],
        ["5", "Dokumen manual penggunaan bagi pengguna, administrator, dan dokumentasi teknis aplikasi",
         "3 dokumen", "Diserahkan dalam format elektronik (PDF/DOCX) dan/atau halaman bantuan daring."],
        ["6", "Pelaksanaan pelatihan/sosialisasi bagi administrator dan perwakilan perangkat daerah",
         "1 kegiatan", "Disertai daftar hadir, materi, dan dokumentasi kegiatan."],
        ["7", "Kode sumber (source code) aplikasi, repositori, dan hak akses pengelolaan",
         "1 paket", "Diserahkan kepada Diskominfo beserta berita acara serah terima."],
        ["8", "Laporan pendahuluan, laporan antara (opsional), laporan akhir, dan berita acara",
         "1 paket", "Disusun sesuai format dan jadwal yang ditetapkan dalam KAK ini."],
    ],
    widths=[1.0, 5.1, 2.0, 6.9],
    align_center_cols=(0,),
)

# =====================================================================
# IV. RUANG LINGKUP & SPESIFIKASI
# =====================================================================
bab("IV", "Ruang Lingkup Pekerjaan dan Spesifikasi Teknis")

sub("A", "Lingkup Pekerjaan")
p("Penyedia jasa wajib melaksanakan pekerjaan pengembangan aplikasi dengan lingkup sebagai berikut:")
bullets([
    "melakukan analisis kebutuhan dan penyusunan rancangan teknis aplikasi berdasarkan hasil "
    "identifikasi kebutuhan Inspektorat dan Dinas Komunikasi dan Informatika;",
    "mengembangkan dan menyempurnakan modul formulir pengisian beserta validasi data;",
    "mengembangkan dan menyempurnakan modul administrator beserta fitur rekapitulasi, pemantauan "
    "status, dan ekspor data;",
    "mengembangkan fitur tanda tangan digital, penyimpanan sementara isian, dan ekspor/cetak dokumen "
    "dalam format PDF ukuran A4;",
    "melakukan penguatan keamanan aplikasi dan penerapan prinsip pelindungan data pribadi;",
    "menyiapkan dan memasang aplikasi pada infrastruktur produksi (hosting, basis data, domain, dan "
    "sertifikat keamanan) serta melakukan migrasi data dari versi awal;",
    "melaksanakan uji fungsi, uji keamanan dasar, dan uji penerimaan pengguna bersama tim teknis dan "
    "perwakilan pengguna;",
    "menyusun dokumentasi teknis dan manual penggunaan, serta melaksanakan pelatihan/sosialisasi;",
    "melaksanakan pendampingan pada masa go-live, pemeliharaan, dan dukungan teknis sesuai jangka "
    "waktu yang ditetapkan dalam kontrak.",
], num=True)
catatan("Aplikasi versi awal beserta kode sumbernya telah tersedia dan dapat diserahkan kepada "
        "penyedia jasa terpilih sesuai ketentuan pengadaan barang/jasa yang berlaku, untuk "
        "selanjutnya dikembangkan, disempurnakan, dan dipasang pada lingkungan produksi.")

sub("B", "Spesifikasi Fungsional")
p("Aplikasi minimal memiliki fungsi-fungsi sebagai berikut:")
table(
    ["No.", "Modul/Fitur", "Persyaratan Fungsi"],
    [
        ["1", "Halaman beranda",
         "Menyajikan informasi umum aplikasi, tautan menuju formulir, serta panduan singkat penggunaan."],
        ["2", "Formulir Daftar Kepentingan Pribadi",
         "Identitas pegawai (nama, NIP, pangkat/golongan, jabatan, perangkat daerah, unit kerja, "
         "tanggal pengisian); Bagian A Hubungan Keluarga dan Kerabat; Bagian B Hubungan Bisnis dan "
         "Finansial; Bagian C Pekerjaan Lain di Luar Pekerjaan Pokok; Bagian D Jabatan Publik Lain "
         "yang Diemban (Rangkap Jabatan); Bagian E Hubungan atau Afiliasi Lainnya; Bagian F Rencana "
         "Pasca Pensiun atau Pengunduran Diri; serta pernyataan dan tanda tangan."],
        ["3", "Formulir Deklarasi Konflik Kepentingan",
         "Identitas pegawai, identitas atasan pejabat yang dituju, tanggal deklarasi, jenis konflik "
         "kepentingan (aktual/potensial), sumber konflik kepentingan, uraian, dan usulan pengendalian."],
        ["4", "Tabel isian dinamis",
         "Penambahan/penghapusan baris isian sesuai kebutuhan, dengan penyimpanan sementara otomatis "
         "di peramban agar data tidak hilang."],
        ["5", "Tanda tangan digital",
         "Kanvas tanda tangan berbasis mouse/sentuhan, dilengkapi fungsi undo dan hapus, disimpan "
         "sebagai citra dalam basis data dengan pembatasan ukuran."],
        ["6", "Cetak/ekspor dokumen",
         "Menghasilkan dokumen ukuran A4 dengan tata letak menyerupai formulir asli, siap dicetak "
         "dan/atau disimpan sebagai PDF."],
        ["7", "Konfirmasi pengisian",
         "Menampilkan bukti/konfirmasi kepada pengisi bahwa data berhasil direkam, tanpa membuka akses "
         "data milik pengguna lain."],
        ["8", "Login administrator",
         "Autentikasi dengan kata sandi yang disimpan dalam bentuk hash (bcrypt/argon2) dan/atau "
         "mekanisme token, dilengkapi pembatasan percobaan login."],
        ["9", "Rekapitulasi dan pencarian",
         "Daftar seluruh pengisian dan deklarasi beserta penyaringan berdasarkan nama, unit kerja, "
         "status, dan periode."],
        ["10", "Detail lembar isian",
         "Menampilkan isian secara read-only dalam format formulir asli lengkap dengan tanda tangan, "
         "siap dicetak."],
        ["11", "Pengelolaan status dan catatan",
         "Pengubahan status penanganan (BARU/DIPROSES/SELESAI/DITOLAK) beserta catatan administrator."],
        ["12", "Ekspor data",
         "Ekspor data rekapitulasi dalam format CSV/Excel untuk kebutuhan pelaporan dan analisis."],
        ["13", "Log aktivitas (audit trail)",
         "Pencatatan aktivitas penting, antara lain waktu login, waktu pengubahan status, dan waktu "
         "penyimpanan data."],
    ],
    widths=[1.0, 4.2, 10.0],
    align_center_cols=(0,),
)

sub("C", "Spesifikasi Non-Fungsional")
table(
    ["No.", "Aspek", "Persyaratan"],
    [
        ["1", "Platform", "Aplikasi berbasis web (browser-based), responsif pada komputer, tablet, dan "
                          "telepon pintar, serta mendukung pengisian melalui layar sentuh."],
        ["2", "Kompatibilitas", "Berfungsi baik pada peramban Google Chrome, Microsoft Edge, Mozilla "
                                "Firefox, dan Safari versi terkini."],
        ["3", "Kinerja", "Waktu muat halaman utama tidak lebih dari 5 (lima) detik pada jaringan "
                         "standar, serta mampu melayani pengisian secara bersamaan pada masa "
                         "sosialisasi tanpa gangguan berarti."],
        ["4", "Keamanan", "Wajib menggunakan HTTPS; kata sandi disimpan dalam bentuk hash; komunikasi "
                          "data menggunakan token terverifikasi; query basis data menggunakan "
                          "parameter terikat (mencegah SQL injection); input dibersihkan (mencegah "
                          "XSS); akses antarsitus (CORS) dibatasi pada domain resmi; pembatasan "
                          "percobaan login; pembatasan ukuran data tanda tangan."],
        ["5", "Pelindungan data pribadi", "Menerapkan prinsip pengumpulan data seperlunya, pembatasan "
                                          "akses berdasarkan peran, penyimpanan yang aman, serta "
                                          "ketentuan masa retensi dan pemusnahan data sesuai ketentuan "
                                          "pelindungan data pribadi."],
        ["6", "Ketersediaan & pencadangan", "Tersedia mekanisme pencadangan (backup) otomatis minimal "
                                            "1 (satu) kali sehari dan prosedur pemulihan data."],
        ["7", "Bahasa & tampilan", "Seluruh antarmuka menggunakan Bahasa Indonesia dengan tata letak "
                                   "yang mengikuti format formulir resmi instansi."],
        ["8", "Interoperabilitas", "Struktur data menggunakan format terbuka (JSON/CSV) dan siap "
                                   "diintegrasikan dengan sistem lain milik Pemerintah Daerah pada "
                                   "tahap pengembangan berikutnya."],
    ],
    widths=[1.0, 4.0, 10.2],
    align_center_cols=(0,),
)

sub("D", "Infrastruktur dan Kebutuhan Sumber Daya")
p("Penyedia jasa wajib menyediakan atau menyiapkan komponen infrastruktur berikut, dan seluruh biaya "
  "termasuk dalam nilai kontrak:")
bullets([
    ("Lingkungan produksi: ", "layanan hosting aplikasi dan basis data relasional (PostgreSQL/MySQL), "
     "termasuk kapasitas penyimpanan yang memadai untuk data isian dan citra tanda tangan."),
    ("Nama domain: ", "subdomain resmi Pemerintah Kabupaten Aceh Tengah atau domain tersendiri yang "
     "dikendalikan oleh Diskominfo."),
    ("Sertifikat keamanan (TLS/SSL): ", "wajib aktif untuk menjamin komunikasi data yang terenkripsi."),
    ("Akun layanan: ", "akun administrator aplikasi, akses pengelolaan basis data, dan akses panel "
     "hosting yang diserahkan kepada Diskominfo pada akhir pekerjaan."),
    ("Kode sumber: ", "diserahkan melalui repositori resmi milik Pemerintah Kabupaten Aceh Tengah "
     "beserta riwayat perubahan (version control)."),
], num=True)

sub("E", "Batasan Pekerjaan")
p("Hal-hal yang tidak termasuk dalam lingkup pekerjaan ini, kecuali diperjanjikan lain secara "
  "tertulis:")
bullets([
    "pengadaan perangkat keras (server, komputer, pemindai, dan perangkat jaringan) serta "
    "pemasangan jaringan;",
    "biaya langganan layanan pihak ketiga di luar komponen sebagaimana butir IV.D;",
    "integrasi menyeluruh dengan sistem kepegawaian atau sistem tanda tangan elektronik tersertifikasi "
    "(dapat dilaksanakan pada tahap pengembangan berikutnya);",
    "pengisian data oleh penyedia atas nama pengguna (kewajiban pengisian tetap berada pada "
    "masing-masing ASN/Pejabat Pemerintahan).",
], num=True)

# =====================================================================
# V. METODE PELAKSANAAN
# =====================================================================
bab("V", "Metode Pelaksanaan")
p("Pelaksanaan pekerjaan dilakukan melalui tahapan berikut:")
table(
    ["No.", "Tahapan", "Uraian Kegiatan", "Keluaran Tahap"],
    [
        ["1", "Persiapan dan analisis kebutuhan",
         "Rapat koordinasi awal, identifikasi kebutuhan pengguna dan administrator, penyusunan "
         "rencana kerja dan jadwal rinci, serta penetapan indikator keberterimaan.",
         "Laporan pendahuluan dan rancangan teknis."],
        ["2", "Perancangan sistem",
         "Penyusunan rancangan antarmuka, struktur dasar data, alur proses bisnis pengisian dan "
         "pengelolaan, serta pemetaan kebutuhan keamanan.",
         "Dokumen rancangan (desain) aplikasi."],
        ["3", "Pengembangan dan penyempurnaan",
         "Pemrograman modul formulir, modul administrator, tanda tangan digital, ekspor dokumen, "
         "penguatan keamanan, serta penyiapan basis data pada lingkungan produksi.",
         "Aplikasi siap uji dan kode sumber."],
        ["4", "Pengujian",
         "Uji fungsi, uji kegunaan, uji keamanan dasar, uji coba pengisian oleh beberapa perwakilan "
         "pengguna, serta perbaikan atas temuan.",
         "Berita acara hasil pengujian."],
        ["5", "Pelatihan dan go-live",
         "Pelatihan/sosialisasi administrator dan perwakilan perangkat daerah, pemasangan aplikasi "
         "pada lingkungan produksi, dan pendampingan penggunaan awal.",
         "Dokumentasi pelatihan dan aplikasi aktif (live)."],
        ["6", "Pemeliharaan dan dukungan teknis",
         "Perbaikan gangguan (bug fixing), pemantauan kinerja dan ketersediaan aplikasi, pencadangan "
         "data, serta dukungan teknis bagi pengguna.",
         "Laporan akhir dan berita acara serah terima."],
    ],
    widths=[1.0, 3.2, 7.8, 3.5],
    align_center_cols=(0,),
)

sub("A", "Ketentuan Pembayaran")
p("Pembayaran dilakukan secara bertahap (termin) sesuai prestasi pekerjaan yang dibuktikan dengan "
  "berita acara, dengan ketentuan sebagai berikut:")
bullets([
    ("Termin I (30%): ", "setelah penandatanganan kontrak dan penyerahan laporan pendahuluan serta "
     "rancangan teknis yang disetujui."),
    ("Termin II (40%): ", "setelah aplikasi selesai dikembangkan dan dinyatakan lulus uji fungsi "
     "berdasarkan berita acara pengujian."),
    ("Termin III (30%): ", "setelah aplikasi aktif pada lingkungan produksi, pelatihan terlaksana, "
     "dan seluruh dokumen serta kode sumber diserahkan melalui berita acara serah terima."),
], num=True)
catatan("Nilai pembayaran yang diterima penyedia merupakan nilai setelah diperhitungkan pajak dan "
        "pungutan lain sesuai ketentuan peraturan perundang-undangan. Mekanisme pembayaran dapat "
        "disesuaikan dengan ketentuan pengadaan barang/jasa yang berlaku pada Pemerintah Kabupaten "
        "Aceh Tengah.")

# =====================================================================
# VI. PELAKSANA DAN PENANGGUNG JAWAB
# =====================================================================
bab("VI", "Pelaksana dan Penanggung Jawab")
table(
    ["No.", "Pihak", "Peran dan Tanggung Jawab"],
    [
        ["1", "Kepala Dinas Komunikasi dan Informatika (selaku Pengguna Anggaran/Kuasa Pengguna "
              "Anggaran)", "Menetapkan kebijakan pelaksanaan kegiatan dan bertanggung jawab atas "
              "penggunaan anggaran."],
        ["2", "Pejabat Penatausahaan Keuangan dan Pejabat Pelaksana Teknis Kegiatan (PPTK)",
         "Melaksanakan pengadaan, mengendalikan pelaksanaan pekerjaan, memverifikasi hasil pekerjaan, "
         "dan menyusun laporan pertanggungjawaban."],
        ["3", "Tim Teknis Diskominfo Kabupaten Aceh Tengah",
         "Memberikan masukan teknis, melakukan pengujian, serta mengawasi kesesuaian hasil pekerjaan "
         "dengan spesifikasi."],
        ["4", "Inspektorat Kabupaten Aceh Tengah",
         "Menyampaikan kebutuhan pengguna, memverifikasi kesesuaian substansi formulir dan alur "
         "pengelolaan konflik kepentingan, serta menjadi pengguna utama aplikasi."],
        ["5", "Penyedia Jasa/Pengembang",
         "Melaksanakan seluruh lingkup pekerjaan sebagaimana KAK ini, menyediakan sumber daya, "
         "menyampaikan laporan, dan menyerahkan hasil pekerjaan secara lengkap."],
        ["6", "Perangkat Daerah Pengguna",
         "Menggunakan aplikasi, mengisi formulir sesuai kewajiban, dan menyampaikan umpan balik atas "
         "kinerja aplikasi."],
    ],
    widths=[1.0, 5.0, 9.3],
    align_center_cols=(0,),
)

# =====================================================================
# VII. JADWAL PELAKSANAAN
# =====================================================================
bab("VII", "Jadwal Pelaksanaan")
p("Jangka waktu pelaksanaan pekerjaan adalah 2 (dua) bulan (8 minggu) sejak tanggal penandatanganan "
  "surat pesanan/kontrak, dengan rincian sebagai berikut:")
table(
    ["No.", "Kegiatan", "M-1", "M-2", "M-3", "M-4", "M-5", "M-6", "M-7", "M-8"],
    [
        ["1", "Persiapan dan analisis kebutuhan", "√", "√", "", "", "", "", "", ""],
        ["2", "Perancangan sistem", "", "√", "√", "", "", "", "", ""],
        ["3", "Pengembangan dan penyempurnaan aplikasi", "", "", "√", "√", "√", "√", "", ""],
        ["4", "Pengujian dan perbaikan", "", "", "", "", "", "√", "√", ""],
        ["5", "Pelatihan/sosialisasi dan go-live", "", "", "", "", "", "", "√", "√"],
        ["6", "Pemeliharaan dan dukungan teknis", "", "", "", "", "", "", "√", "√"],
        ["7", "Laporan akhir dan serah terima", "", "", "", "", "", "", "", "√"],
    ],
    widths=[1.0, 4.6, 0.9, 0.9, 0.9, 0.9, 0.9, 0.9, 0.9, 0.9],
    align_center_cols=tuple(range(2, 10)),
    size=9.5,
)
catatan("M-1 s.d. M-8 = minggu ke-1 sampai dengan minggu ke-8 masa pelaksanaan. Jadwal terinci "
        "dituangkan dalam rencana kerja penyedia dan disesuaikan dengan tanggal kontrak. Apabila "
        "terdapat perubahan jadwal, wajib mendapat persetujuan tertulis dari PPTK.")

# =====================================================================
# VIII. RINCIAN BIAYA
# =====================================================================
bab("VIII", "Rincian Biaya")
p("Rincian anggaran pengembangan aplikasi Dek-Oke bersumber dari " + PARAM["sumber_dana"] +
  f" dengan pagu sebesar {PARAM['pagu']} ({PARAM['pagu_kata']}), sebagai berikut:")
_rab = table(
    ["No.", "Uraian Pekerjaan/Barang", "Satuan", "Vol.", "Harga Satuan (Rp)", "Jumlah (Rp)"],
    [
        ["1", "Pengembangan dan penyempurnaan modul formulir (Daftar Kepentingan Pribadi Bagian A–F "
              "dan Deklarasi Konflik Kepentingan) termasuk validasi dan penyimpanan sementara",
         "paket", "1", "2.750.000", "2.750.000"],
        ["2", "Pengembangan modul administrator (rekapitulasi, detail lembar isian siap cetak, "
              "pengelolaan status dan catatan, ekspor data)",
         "paket", "1", "1.750.000", "1.750.000"],
        ["3", "Pengembangan fitur tanda tangan digital serta ekspor/cetak dokumen PDF ukuran A4",
         "paket", "1", "1.000.000", "1.000.000"],
        ["4", "Penguatan keamanan aplikasi, pelindungan data pribadi, dan log aktivitas",
         "paket", "1", "1.000.000", "1.000.000"],
        ["5", "Sewa hosting aplikasi dan basis data, domain/subdomain, serta sertifikat keamanan "
              "(12 bulan)", "bulan", "12", "125.000", "1.500.000"],
        ["6", "Pelatihan/sosialisasi administrator dan perwakilan perangkat daerah serta pendampingan "
              "go-live", "kegiatan", "1", "1.000.000", "1.000.000"],
        ["7", "Penyusunan manual penggunaan, dokumentasi teknis, dan laporan akhir",
         "paket", "1", "500.000", "500.000"],
        ["8", "Dukungan teknis dan pencadangan data pascaimplementasi (2 bulan)",
         "bulan", "2", "250.000", "500.000"],
        ["", "JUMLAH", "", "", "", "10.000.000"],
    ],
    widths=[1.0, 7.4, 1.5, 1.0, 2.3, 2.3],
    align_center_cols=(0, 2, 3),
    align_right_cols=(4, 5),
    size=10,
)
# Baris JUMLAH: gabungkan kolom uraian dan buat tebal
_baris = _rab.rows[-1]
_gabung = _baris.cells[0].merge(_baris.cells[4])
_gabung.text = ""
_par = _gabung.paragraphs[0]
_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
_par.paragraph_format.space_after = Pt(2)
_par.add_run("JUMLAH").bold = True
_fix_font(_par, size=10, bold=True)
for _r in _baris.cells[5].paragraphs[0].runs:
    _r.bold = True
p_mixed([("Terbilang: ", True), ("Sepuluh Juta Rupiah", False)], space_after=4)
p("Ketentuan biaya:")
bullets([
    "nilai tersebut merupakan pagu maksimum dan sudah termasuk seluruh biaya komponen pekerjaan "
    "sebagaimana butir IV.D;",
    "pajak dan pungutan lain dibebankan sesuai ketentuan peraturan perundang-undangan;",
    "penawaran yang melebihi pagu dinyatakan tidak sah;",
    "sisa anggaran (bila ada) tidak dapat digunakan di luar kontrak dan menjadi penghematan anggaran "
    "kegiatan.",
], num=True)

# =====================================================================
# IX. INDIKATOR KINERJA
# =====================================================================
bab("IX", "Indikator Kinerja dan Manfaat")
table(
    ["No.", "Indikator Kinerja", "Target", "Cara Pengukuran"],
    [
        ["1", "Ketersediaan aplikasi pada lingkungan produksi", "100%",
         "Pemeriksaan langsung dan pengujian akses oleh tim teknis."],
        ["2", "Jumlah modul formulir yang berfungsi sesuai kebutuhan",
         "2 modul (Daftar Kepentingan Pribadi dan Deklarasi Konflik Kepentingan)",
         "Berita acara uji fungsi."],
        ["3", "Tingkat keberhasilan penyimpanan data pengisian", "≥ 99%",
         "Uji coba pengisian dan pemeriksaan basis data."],
        ["4", "Pemenuhan aspek keamanan dasar", "100% terpenuhi",
         "Daftar periksa keamanan (checklist) dan hasil pengujian."],
        ["5", "Pelaksanaan pelatihan/sosialisasi", "1 kegiatan, kehadiran ≥ 80% undangan",
         "Daftar hadir dan dokumentasi kegiatan."],
        ["6", "Penyerahan dokumentasi dan kode sumber", "100% lengkap",
         "Berita acara serah terima."],
        ["7", "Kepuasan pengguna awal terhadap kemudahan penggunaan",
         "≥ 80% responden menyatakan mudah",
         "Kuesioner umpan balik pengguna."],
    ],
    widths=[1.0, 4.6, 4.4, 5.3],
    align_center_cols=(0,),
)

# =====================================================================
# X. PELAPORAN, MONITORING, DAN EVALUASI
# =====================================================================
bab("X", "Pelaporan, Monitoring, dan Evaluasi")
sub("A", "Pelaporan")
bullets([
    ("Laporan Pendahuluan: ", "diserahkan paling lambat 7 (tujuh) hari kerja setelah penandatanganan "
     "kontrak, memuat rencana kerja, jadwal rinci, dan rancangan teknis awal."),
    ("Laporan Antara (bila diperlukan): ", "diserahkan pada saat pengujian, memuat kemajuan "
     "pekerjaan, kendala, dan rencana tindak lanjut."),
    ("Laporan Akhir: ", "diserahkan paling lambat 7 (tujuh) hari kerja setelah seluruh pekerjaan "
     "selesai, memuat uraian pelaksanaan, hasil pengujian, dokumentasi, dan lampiran berita acara."),
    ("Berita Acara: ", "berita acara pengujian, berita acara serah terima aplikasi (termasuk kode "
     "sumber dan akun layanan), serta berita acara penyelesaian pekerjaan."),
], num=True)

sub("B", "Monitoring dan Evaluasi")
bullets([
    "PPTK bersama tim teknis melaksanakan pemantauan kemajuan pekerjaan secara berkala, sekurang-"
    "kurangnya 1 (satu) kali dalam 2 (dua) minggu, dan menuangkan hasilnya dalam catatan pemantauan.",
    "Pengujian penerimaan dilakukan bersama Inspektorat dan perwakilan pengguna sebelum aplikasi "
    "dinyatakan aktif (go-live).",
    "Evaluasi pemanfaatan aplikasi dilakukan paling lambat 3 (tiga) bulan setelah go-live, meliputi "
    "jumlah pengguna, jumlah isian, kendala teknis, dan rencana pengembangan lanjutan.",
    "Hasil evaluasi menjadi dasar penyusunan rencana pengembangan dan pemeliharaan pada tahun "
    "anggaran berikutnya.",
], num=True)

# =====================================================================
# XI. KEPEMILIKAN HASIL, KERAHASIAAN, PELINDUNGAN DATA
# =====================================================================
bab("XI", "Kepemilikan Hasil Pekerjaan, Kerahasiaan, dan Pelindungan Data")
bullets([
    ("Kepemilikan hasil pekerjaan: ", "seluruh hasil pekerjaan, termasuk kode sumber, basis data, "
     "dokumentasi, desain, dan hak kekayaan intelektual atas aplikasi, menjadi milik Pemerintah "
     "Kabupaten Aceh Tengah dan tidak boleh digunakan, dijual, atau dialihkan kepada pihak lain "
     "tanpa persetujuan tertulis."),
    ("Serah terima akses: ", "penyedia wajib menyerahkan seluruh akun, kata sandi, kunci akses, dan "
     "hak pengelolaan infrastruktur pada saat serah terima, serta menghapus salinan data dari "
     "perangkat miliknya."),
    ("Kerahasiaan data: ", "data yang dikelola bersifat terbatas dan memuat data pribadi pegawai. "
     "Penyedia dan seluruh personelnya wajib menjaga kerahasiaan data dan dilarang mengakses, "
     "menggandakan, atau menyebarluaskan data di luar keperluan pekerjaan."),
    ("Pelindungan data pribadi: ", "pengelolaan data pribadi dilaksanakan sesuai Undang-Undang Nomor "
     "27 Tahun 2022 tentang Pelindungan Data Pribadi, meliputi pembatasan tujuan pengumpulan, "
     "pembatasan akses, keamanan penyimpanan, serta ketentuan retensi dan pemusnahan data."),
    ("Sanksi: ", "pelanggaran atas ketentuan kerahasiaan dan pelindungan data dikenakan sanksi sesuai "
     "ketentuan peraturan perundang-undangan dan/atau ketentuan dalam kontrak."),
], num=True)

# =====================================================================
# XII. PENUTUP
# =====================================================================
bab("XII", "Penutup")
p("Kerangka Acuan Kerja ini disusun sebagai acuan dalam pelaksanaan pengembangan aplikasi Dek-Oke "
  "(Deklarasi Konflik Kepentingan) pada " + PARAM["opd_pelaksana"] + " Tahun Anggaran " +
  PARAM["tahun"] + ". Hal-hal yang belum diatur atau memerlukan penyesuaian dalam pelaksanaan "
  "kegiatan akan ditetapkan kemudian melalui kesepakatan para pihak dengan tetap mengacu pada "
  "ketentuan peraturan perundang-undangan yang berlaku.")
p("Demikian Kerangka Acuan Kerja ini dibuat untuk dipergunakan sebagaimana mestinya.")

# =====================================================================
# LAMPIRAN
# =====================================================================
doc.add_page_break()
p("LAMPIRAN 1", size=12, bold=True, align="center", space_after=2)
p("DAFTAR PERIKSA KRITERIA KEBERTERIMAAN (UJI TERIMA)", size=12, bold=True, align="center", space_after=10)
table(
    ["No.", "Aspek", "Kriteria Keberterimaan", "Sesuai", "Tidak"],
    [
        ["1", "Fungsi formulir", "Formulir Daftar Kepentingan Pribadi (identitas dan Bagian A–F) "
         "dapat diisi, divalidasi, dan disimpan tanpa kesalahan.", "☐", "☐"],
        ["2", "Fungsi formulir", "Formulir Deklarasi Konflik Kepentingan dapat diisi dan disimpan "
         "lengkap dengan identitas atasan pejabat.", "☐", "☐"],
        ["3", "Tanda tangan", "Tanda tangan digital tersimpan, tampil kembali dengan benar pada lembar "
         "isian, dan tercetak pada dokumen.", "☐", "☐"],
        ["4", "Dokumen", "Hasil cetak/ekspor PDF berukuran A4 dengan tata letak menyerupai formulir "
         "asli dan dapat dibaca dengan baik.", "☐", "☐"],
        ["5", "Administrator", "Login, rekapitulasi, pencarian, pengubahan status dan catatan, serta "
         "ekspor data berfungsi dengan benar dan hanya dapat diakses oleh pengguna berwenang.", "☐", "☐"],
        ["6", "Keamanan", "Seluruh akses menggunakan HTTPS; kata sandi tersimpan dalam bentuk hash; "
         "percobaan login dibatasi; tidak ditemukan kerentanan dasar (SQL injection/XSS).", "☐", "☐"],
        ["7", "Data", "Basis data produksi terpasang; pencadangan otomatis harian berjalan dan "
         "prosedur pemulihan telah diuji.", "☐", "☐"],
        ["8", "Kinerja", "Halaman utama termuat ≤ 5 detik; pengisian bersamaan pada masa sosialisasi "
         "berjalan tanpa gangguan berarti.", "☐", "☐"],
        ["9", "Dokumentasi", "Manual pengguna, manual administrator, dan dokumentasi teknis lengkap "
         "serta mudah dipahami.", "☐", "☐"],
        ["10", "Alih pengetahuan", "Pelatihan terlaksana, kode sumber dan seluruh akun/akses telah "
         "diserahkan melalui berita acara.", "☐", "☐"],
    ],
    widths=[1.0, 2.6, 8.8, 1.5, 1.5],
    align_center_cols=(0, 3, 4),
)

doc.add_page_break()
p("LAMPIRAN 2", size=12, bold=True, align="center", space_after=2)
p("STRUKTUR DATA MINIMAL APLIKASI", size=12, bold=True, align="center", space_after=10)
p("Struktur basis data minimal yang harus tersedia pada lingkungan produksi adalah sebagai berikut:")
table(
    ["No.", "Tabel", "Kolom Utama", "Keterangan"],
    [
        ["1", "pengisian (Daftar Kepentingan Pribadi)",
         "id; nama; nip; pangkat; jabatan; perangkat; unit_kerja; tanggal_isi; bagian_a s.d. bagian_f; "
         "ttd; ttd_nama; ttd_nip; status; catatan_admin; created_at",
         "Bagian A–F disimpan dalam format terstruktur (JSON) agar jumlah baris isian fleksibel."],
        ["2", "deklarasi_konflik (Deklarasi Konflik Kepentingan)",
         "id; nama; nip; jabatan; unit_kerja; perangkat; atasan_nama; atasan_nip; atasan_jabatan; "
         "atasan_unit; atasan_perangkat; tanggal_isi; jenis_konflik; sumber_konflik; uraian; "
         "pengendalian; ttd; ttd_nama; ttd_nip; status; catatan_admin; created_at",
         "Menyimpan deklarasi beserta identitas atasan pejabat yang dituju."],
        ["3", "log_aktivitas (opsional)",
         "id; pengguna; aksi; keterangan; waktu",
         "Mencatat aktivitas penting untuk keperluan penelusuran (audit trail)."],
    ],
    widths=[1.0, 3.8, 6.6, 4.0],
    align_center_cols=(0,),
    size=10,
)

doc.add_page_break()
p("LAMPIRAN 3", size=12, bold=True, align="center", space_after=2)
p("ALUR PROSES PENGGUNAAN APLIKASI", size=12, bold=True, align="center", space_after=10)
table(
    ["No.", "Tahap", "Pelaku", "Uraian"],
    [
        ["1", "Penyiapan aplikasi", "Diskominfo dan Penyedia",
         "Aplikasi dipasang pada lingkungan produksi, basis data disiapkan, akun administrator "
         "dibuat, dan akses diuji."],
        ["2", "Sosialisasi dan pelatihan", "Diskominfo dan Inspektorat",
         "Sosialisasi kepada perangkat daerah mengenai kewajiban pengisian dan tata cara penggunaan "
         "aplikasi."],
        ["3", "Pengisian formulir", "ASN/Pejabat Pemerintahan",
         "Mengakses aplikasi, mengisi formulir, membubuhkan tanda tangan digital, dan mengirim data."],
        ["4", "Perekaman dan konfirmasi", "Aplikasi",
         "Data tersimpan pada basis data; pengisi menerima konfirmasi bahwa data berhasil direkam."],
        ["5", "Pemantauan dan verifikasi", "Inspektorat (Administrator)",
         "Memeriksa daftar isian, menelusuri detail lembar isian, serta menetapkan status dan catatan "
         "penanganan."],
        ["6", "Analisis dan pengendalian", "Inspektorat dan Atasan Pejabat",
         "Menindaklanjuti deklarasi konflik kepentingan sesuai bentuk pengendalian yang ditetapkan "
         "(misalnya pengalihan tugas atau penggantian pengambil keputusan)."],
        ["7", "Pelaporan", "Inspektorat",
         "Menyusun rekapitulasi dan laporan pemanfaatan aplikasi sebagai bahan evaluasi dan "
         "pertanggungjawaban."],
        ["8", "Pemeliharaan dan pengembangan", "Diskominfo dan Penyedia",
         "Melakukan pemeliharaan rutin, pencadangan data, perbaikan gangguan, dan rencana "
         "pengembangan lanjutan."],
    ],
    widths=[1.0, 3.6, 3.6, 7.2],
    align_center_cols=(0,),
)

# --- penutup lampiran
p()
p("Dokumen ini merupakan bagian yang tidak terpisahkan dari Kerangka Acuan Kerja "
  "Pengembangan Aplikasi Dek-Oke Tahun Anggaran " + PARAM["tahun"] + ".", italic=True,
  align="center", size=11)

doc.save(OUT)
print("Berhasil dibuat:", OUT)
